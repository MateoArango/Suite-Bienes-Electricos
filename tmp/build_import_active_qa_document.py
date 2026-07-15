from __future__ import annotations

import json
import re
import sys
from pathlib import Path

TOOLS = Path(r"C:\Users\marango\Desktop\Suite Bienes Eléctricos\.tools")
sys.path.insert(0, str(TOOLS))

from docx import Document  # type: ignore
from docx.enum.section import WD_ORIENT  # type: ignore
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT  # type: ignore
from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
from docx.oxml import OxmlElement  # type: ignore
from docx.oxml.ns import qn  # type: ignore
from docx.shared import Inches, Pt, RGBColor  # type: ignore


ROOT = Path(r"C:\Users\marango\Desktop\Suite Bienes Eléctricos")
REFERENCE_JSON = ROOT / "tmp" / "reference_extracts" / "reference_summary.json"
OUTPUT = ROOT / "output" / "docx" / "ImportActiveComponent_Automation_Test_Scenario_Map.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

reference = json.loads(REFERENCE_JSON.read_text(encoding="utf-8"))
plantilla = next(sheet for sheet in reference["xlsx"]["sheets"] if sheet["title"] == "Plantilla")

# compact_reference_guide, with a named landscape override for a scenario-map document.
FONT = "Calibri"
MONO = "Consolas"
NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "F5F8FB"
GREEN = "E7F4EC"
GREEN_TEXT = "23633C"
AMBER = "FFF4D6"
AMBER_TEXT = "7A5A00"
RED = "FBE9E9"
RED_TEXT = "9B1C1C"
WHITE = "FFFFFF"
BORDER = "AEB7C2"
BLACK = "111827"

PAGE_W = 11.0
PAGE_H = 8.5
MARGIN = 0.5
CONTENT_W = 10.0
CONTENT_DXA = 14400
TABLE_INDENT_DXA = 120


def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=8.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in, indent_dxa=TABLE_INDENT_DXA):
    widths_dxa = [round(value * 1440) for value in widths_in]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_table_borders(table, color=BORDER, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths, *, font_size=8.0, priority_column=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        set_cell_shading(header.cells[index], LIGHT_BLUE)
        set_cell_text(header.cells[index], label, bold=True, color=NAVY, size=8.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 1, len(headers) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[index], value, size=font_size, align=align)
        if priority_column is not None:
            priority = str(row_data[priority_column])
            if priority == "P0":
                set_cell_shading(row.cells[priority_column], RED)
                for run in row.cells[priority_column].paragraphs[0].runs:
                    set_font(run, size=font_size, bold=True, color=RED_TEXT)
            elif priority == "P1":
                set_cell_shading(row.cells[priority_column], AMBER)
                for run in row.cells[priority_column].paragraphs[0].runs:
                    set_font(run, size=font_size, bold=True, color=AMBER_TEXT)
            else:
                set_cell_shading(row.cells[priority_column], GREEN)
                for run in row.cells[priority_column].paragraphs[0].runs:
                    set_font(run, size=font_size, bold=True, color=GREEN_TEXT)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    return p


def add_body(doc, text, *, bold_lead=None, after=5, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, size=9.4, bold=True, color=BLACK)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest, size=9.4, italic=italic, color=BLACK)
    else:
        run = p.add_run(text)
        set_font(run, size=9.4, italic=italic, color=BLACK)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_font(run, size=9.2, color=BLACK)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W], indent_dxa=TABLE_INDENT_DXA)
    set_table_borders(table, color=BORDER, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.12
    lead = p.add_run(f"{label}: ")
    set_font(lead, size=9.3, bold=True, color=color)
    body = p.add_run(text)
    set_font(body, size=9.3, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W], indent_dxa=TABLE_INDENT_DXA)
    set_table_borders(table, color="D6DCE3", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F8FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code)
    set_font(run, name=MONO, size=7.6, color="263445")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_field(run, field_code):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def clean_comment(comment):
    if not comment:
        return "No explicit comment constraint"
    text = re.sub(r"^[^A-Za-zÁÉÍÓÚáéíóúÑñ0-9]+\s*", "", comment.strip())
    return " | ".join(part.strip() for part in text.splitlines() if part.strip())


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(PAGE_W)
section.page_height = Inches(PAGE_H)
section.top_margin = Inches(MARGIN)
section.bottom_margin = Inches(MARGIN)
section.left_margin = Inches(MARGIN)
section.right_margin = Inches(MARGIN)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(9.4)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 14, 7),
    ("Heading 2", 13, BLUE, 11, 6),
    ("Heading 3", 11, DARK_BLUE, 8, 4),
):
    style = styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for style_name, left, first in (("List Bullet", 0.38, -0.19), ("List Bullet 2", 0.7, -0.18)):
    style = styles[style_name]
    style.font.name = FONT
    style.font.size = Pt(9.2)
    style.paragraph_format.left_indent = Inches(left)
    style.paragraph_format.first_line_indent = Inches(first)
    style.paragraph_format.space_after = Pt(3)

# Running header and footer.
header = section.header
header_table = header.add_table(rows=1, cols=2, width=Inches(CONTENT_W))
set_table_geometry(header_table, [5.0, 5.0], indent_dxa=0)
for cell in header_table.rows[0].cells:
    set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
left_p = header_table.cell(0, 0).paragraphs[0]
left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
left_r = left_p.add_run("ImportActiveComponent")
set_font(left_r, size=8.2, bold=True, color=MUTED)
right_p = header_table.cell(0, 1).paragraphs[0]
right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
right_r = right_p.add_run("QA automation scenario map | 2026-07-14")
set_font(right_r, size=8.2, color=MUTED)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_run = footer_p.add_run("Electrical Assets Import | Page ")
set_font(footer_run, size=8.0, color=MUTED)
field_run = footer_p.add_run()
set_font(field_run, size=8.0, color=MUTED)
add_field(field_run, "PAGE")

# First-page memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("QA AUTOMATION TEST STRATEGY")
set_font(r, size=10.5, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(3)
r = p.add_run("ImportActiveComponent - Comprehensive Test Scenario Map")
set_font(r, size=24, bold=True, color=NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Playwright + APIRequestContext + Excel contract checks with SheetJS")
set_font(r, size=12.5, color=MUTED)

meta_rows = [
    ("Route", "/dashboard/bienelectrico/importar"),
    ("UI context", "Dashboard panel, 600 px wide"),
    ("Backend load type", "1 | description=importacion | authenticated login attached"),
    ("Evidence reviewed", "Current Angular component/services, supplied XLSX, test-id reference, edit/create QA maps, Playwright suite"),
]
add_table(doc, ["Field", "Value"], meta_rows, [1.65, 8.35], font_size=8.7)

add_callout(
    doc,
    "Recommendation",
    "Automate the import as a state-changing electrical-asset workflow, not as a generic uploader. Keep deep end-to-end coverage on template compatibility, validation, final POST behavior, and post-import GET reconciliation. Give History and error-CSV export only one stable contract test each because their business value is low.",
    fill=LIGHT_BLUE,
)

add_heading(doc, "1. What the evidence says", 1)
evidence_rows = [
    ("Workbook", "3 sheets: Plantilla, Instrucciones, Codigos Dane; 102 headers from A to CX."),
    ("Required fields", "7 dark-header columns: ARTICULO, N° PLACA, FOTOS (enlace), PLANILLA (ARCGIS), DEPARTAMENTO, MUNICIPIO, COD_LOCALIZACION_DANE."),
    ("Optional fields", "95 light-green columns. Header comments carry type and maximum-length guidance."),
    ("Location rules", "Departamento and Municipio use dependent validation lists through row 5000. Instructions say DANE codes are persisted."),
    ("Formatting rules", "Decimal point, dates DD/MM/AAAA, and complete https:// URLs are explicitly documented in the workbook."),
    ("Frontend gate", ".xlsx/.xls, one file, maximum 10 MB; successful backend validation is required before import."),
    ("Error handling", "HTTP 400, 402, and 422 open the detailed validation-error path; other statuses use snackbar/transient error handling."),
    ("Current locator risk", "The supplied PDF lists 9 data-testid values, but the inspected HTML snapshot contains none. Confirm the target branch before writing locators."),
]
add_table(doc, ["Area", "Observed contract"], evidence_rows, [1.7, 8.3], font_size=8.5)

add_heading(doc, "2. Automation principles", 1)
for text in (
    "Use the downloaded backend template as the source of truth. A version-controlled header snapshot should fail loudly when the backend changes the column contract.",
    "Acquire available plates dynamically from GET /electrical-assets/article-resume/available-plates. Do not hard-code plates that other tests can consume.",
    "Treat each successful import as destructive test data. Run state-changing import suites serially and reserve disposable QA plates.",
    "Assert both network and persistence: validate POST, final load POST, then GET /electrical-assets/article-resume/{plate} for the database-facing result.",
    "Separate deterministic UI tests with mocked responses from a small number of live backend contract/E2E tests.",
    "For update and blank-clearing behavior, restore the original record in finally when the API supports a reliable PATCH restore; otherwise use resettable seeded data only.",
):
    add_bullet(doc, text)

add_heading(doc, "3. Recommended execution portfolio", 1)
portfolio_rows = [
    ("Pull request", "P0 local gates and mocked API/UI states; template header contract; no bulk state consumption.", "Fast and deterministic"),
    ("Daily smoke", "One disposable valid plate: validate, import, post-import GET, cleanup/reset.", "Serial worker"),
    ("Nightly", "Existing-plate replacement, blank clearing, 25-50 valid plates, atomicity, retry, pagination.", "Dedicated seed pool"),
    ("Release", "10 MB boundary, .xls compatibility, 4999-row workbook boundary, concurrency and performance.", "Controlled environment"),
]
add_table(doc, ["Cadence", "Scope", "Data rule"], portfolio_rows, [1.35, 6.25, 2.4], font_size=8.5)

add_heading(doc, "4. Comprehensive scenario matrix", 1)
add_body(doc, "Priority is risk-based: P0 protects creation/update integrity, P1 protects validation and recovery, and P2 covers secondary UX. A scenario marked Contract can run below the UI; UI+API means Playwright drives the component and captures/reconciles network behavior.")

scenario_headers = ["ID", "Pri", "Layer", "Scenario", "Test data", "Key assertions", "Cadence"]
scenario_widths = [0.82, 0.42, 0.78, 1.62, 1.52, 3.98, 0.86]

suites = [
    (
        "4.1 Template, access, and local uploader gates",
        [
            ("QA-IMP-001", "P0", "UI", "Route and authentication gate", "Authenticated and unauthenticated sessions", "Authenticated user can open the 600 px panel. Unauthenticated/expired session is redirected or denied. Import tab is default and submit starts disabled.", "PR"),
            ("QA-IMP-002", "P0", "UI+API", "Download current template", "Backend template blob", "Exactly one GET /import/template?loadTypeId=1; response is non-empty Excel; downloaded name matches plantilla_importacion_YYYY-MM-DD_HH.mm.xlsx.", "PR"),
            ("QA-IMP-003", "P0", "Contract", "SheetJS workbook/header contract", "Downloaded template", "XLSX.read succeeds; sheets are Plantilla/Instrucciones/Codigos Dane; first row has exactly 102 unique headers in approved order from A to CX.", "PR"),
            ("QA-IMP-004", "P0", "Contract", "Required-column contract", "Downloaded template", "Approved required set is exactly A,B,C,D,E,F,Z. Comments/types remain available. Flag a backend template change instead of silently adapting the test.", "PR"),
            ("QA-IMP-005", "P1", "UI+API", "Accept valid .xlsx", "valid_minimal_1.xlsx", "File is readable, validation POST occurs once, uploader reaches completed, and Importar activo becomes enabled only after validation success.", "PR"),
            ("QA-IMP-006", "P1", "UI+API", "Accept real legacy .xls", "valid_minimal_1.xls", "A genuine BIFF .xls is accepted and validated. Do not satisfy this with an .xlsx merely renamed .xls.", "Release"),
            ("QA-IMP-007", "P1", "UI", "Reject local file misuse", "pdf/csv/txt, renamed file, unreadable UploadedFile, two-file drop", "Unsupported extension, unreadable file, and multiple-file attempts never call validation; uploader shows the correct local error and submit remains disabled.", "PR"),
            ("QA-IMP-008", "P1", "UI", "10 MB boundary", "10 MB exactly; 10 MB + 1 byte", "A file at the limit reaches backend validation; one byte over is rejected locally as Archivo muy grande; no validation request for the oversized file.", "Release"),
        ],
    ),
    (
        "4.2 Backend validation and spreadsheet business rules",
        [
            ("QA-IMP-009", "P0", "UI+API", "Header-only or empty workbook", "header_only.xlsx; truly blank workbook", "Backend response with totalRows=0 yields El archivo no contiene datos, uploader Sin datos, null validatedFile, and disabled submit.", "PR"),
            ("QA-IMP-010", "P1", "Contract", "Sparse rows", "Rows 2 and 4 populated; rows 3/5 blank", "Blank lines are ignored consistently; totalRows counts actual asset rows; error row numbers still point to physical Excel rows.", "Nightly"),
            ("QA-IMP-011", "P0", "Contract", "Each required cell empty", "7 generated files, one missing required value per file", "Validation rejects ARTICULO, N° PLACA, FOTOS, PLANILLA, DEPARTAMENTO, MUNICIPIO, and COD_LOCALIZACION_DANE independently with row, column, description.", "PR"),
            ("QA-IMP-012", "P0", "Contract", "Header mutation matrix", "missing, renamed, duplicate, reordered, extra, whitespace header", "Backend behavior is explicit for every mutation. Import cannot proceed on a structurally incompatible workbook. Reordering is asserted according to the agreed backend contract, not guessed.", "PR"),
            ("QA-IMP-013", "P0", "API+UI", "Valid available plate", "Plate reserved from available-plates API", "Validation succeeds only when the plate/article relationship is eligible. The UI enables import and preserves the exact File object for final load.", "Daily"),
            ("QA-IMP-014", "P0", "Contract", "Unknown or unavailable plate", "Unknown, baja/devolución, already consumed plate", "Validation rejects ineligible plates with an actionable row/column error; no final load request is possible.", "Nightly"),
            ("QA-IMP-015", "P0", "Contract", "Duplicate plate within workbook", "Same plate twice; conflicting values", "Validation reports both the duplication and relevant row numbers. No partial import or ambiguous last-row-wins behavior occurs unless explicitly specified.", "Nightly"),
            ("QA-IMP-016", "P0", "Contract", "Location relationship", "Valid/invalid department, municipality and DANE combinations", "Department and municipality are valid codes, municipality belongs to department, and COD_LOCALIZACION_DANE is consistent. Name-only and code-name inputs follow workbook instructions.", "Nightly"),
            ("QA-IMP-017", "P1", "Contract", "Types, lengths, and formats", "Boundary matrix from header comments", "For representative text/number/decimal/date/URL fields: max is accepted, max+1 rejected; comma decimals, impossible dates, malformed URLs, negative/decimal values are handled by documented backend rules.", "Nightly"),
            ("QA-IMP-018", "P1", "Contract", "Whitespace, Unicode, and formulas", "Trimmed values, accents, emoji, formulas, numeric-looking text", "Leading zeros in plate are preserved; whitespace normalization is deterministic; formulas are rejected or evaluated per contract; Spanish accents survive end to end.", "Nightly"),
            ("QA-IMP-019", "P1", "UI+API", "Detailed list errors", "400, 402, 422 with errors/data/data.errors/array/message", "Uploader shows Datos inválidos; summary alert opens; adapters display Spanish or English keys; dialog preserves row/column/description; import stays disabled.", "PR"),
            ("QA-IMP-020", "P1", "UI+API", "Non-list validation failures and recovery", "401, 403, 413, 500, timeout", "Snackbar/transient message is visible; retry triggers a new validation; cancel/remove clears every validation/import signal; a prior success cannot enable a newly failed file.", "PR"),
        ],
    ),
    (
        "4.3 Final import, API contract, and persisted asset state",
        [
            ("QA-IMP-021", "P0", "UI+API", "Create one asset end to end", "valid_minimal_1.xlsx + disposable available plate", "Validation succeeds; one final load POST succeeds; success snackbar appears; uploader clears; GET article-resume/{plate} returns the created asset with expected required values.", "Daily"),
            ("QA-IMP-022", "P0", "API", "Final POST request contract", "Validated workbook", "POST /import/load uses multipart field file with original filename and query params loadTypeId=1, description=importacion, user=current login. Browser supplies multipart boundary.", "PR"),
            ("QA-IMP-023", "P0", "API", "Pre/post state reconciliation", "Reserved available plate", "Capture precondition from available-plates; after load, poll GET article-resume/{plate}; compare normalized fields; confirm plate is no longer offered as available when business rules require it.", "Daily"),
            ("QA-IMP-024", "P0", "API+UI", "Mass import valid plates", "25-50 dynamically reserved plates", "Validation totalRows=N; load succeeds once; every plate is retrievable; no duplicates/missing records; history processed count equals N. Record validation/load duration without a fragile hard SLA.", "Nightly"),
            ("QA-IMP-025", "P0", "API", "Existing plate full replacement", "update_full_existing_plate.xlsx", "GET baseline, import changed values across multiple asset sections, GET again, and compare the complete expected snapshot. Prove import replacement semantics rather than assuming PATCH behavior.", "Nightly"),
            ("QA-IMP-026", "P0", "API", "Blank optional cells clear values", "update_clear_optional.xlsx", "Start with non-empty optional values, leave selected cells blank, import, then GET. Assert the canonical cleared representation (null or empty string) field by field and that required identifiers remain intact.", "Nightly"),
            ("QA-IMP-027", "P0", "API", "Mixed valid/invalid batch atomicity", "One valid + one invalid row", "Validation blocks final load. If load is forced at API level, assert the agreed all-or-nothing behavior and verify the valid row was not partially persisted.", "Nightly"),
            ("QA-IMP-028", "P0", "UI+API", "Double submit and re-import", "Validated workbook; rapid click; same workbook twice", "Only one load request from rapid clicks. Re-importing the same plate/workbook does not create duplicates; response follows the documented update/idempotency contract.", "PR"),
            ("QA-IMP-029", "P1", "UI+API", "Load failure then retry", "500/timeout on first load; success second", "Validated state is retained after final-load failure; message is visible; retry sends the same file once; success clears uploader. Validation is not repeated unless product requires it.", "PR"),
            ("QA-IMP-030", "P1", "API", "Concurrent imports", "Two workbooks sharing one plate; disjoint control", "Same-plate conflict has one deterministic winner/rejection and no corrupt merge. Disjoint imports remain isolated. Run only with resettable fixtures.", "Release"),
            ("QA-IMP-031", "P1", "API", "Authenticated user attribution", "Known QA login", "Final POST user matches authenticated login and the resulting load/history record attributes the operation consistently. Empty/stale login is rejected or surfaced.", "Nightly"),
        ],
    ),
    (
        "4.4 History, error export, layout, and UX",
        [
            ("QA-IMP-032", "P2", "UI+API", "History request contract", "Mocked page 0", "Selecting Historial sends GET /electrical-assets/loads?page=0&size=5&sort=createdAt,desc; newest records render first with date, status, records, description.", "PR"),
            ("QA-IMP-033", "P2", "UI", "Infinite scroll pagination", "3 mocked pages; overlapping IDs", "Near-bottom scroll loads one next page at a time; loading guard prevents duplicates; last=true stops requests; repeated IDs follow the agreed rendering rule.", "Nightly"),
            ("QA-IMP-034", "P2", "UI", "History empty state", "Empty array and empty paginated response", "No hay registros para mostrar is visible and no endless loader remains.", "PR"),
            ("QA-IMP-035", "P2", "UI", "History/template failure feedback", "500/timeout", "Document current gap: failures only stop loading/log to console. Keep one regression test pending until visible feedback is implemented.", "Backlog"),
            ("QA-IMP-036", "P2", "UI+API", "History after successful import", "Successful load then tab switch", "Document current behavior: history is not refreshed immediately but is reloaded when Historial is selected. Assert the actual product decision once confirmed.", "Nightly"),
            ("QA-IMP-037", "P2", "Component", "Validation-error CSV contract", "Spanish/English error keys, separators, quotes, newlines", "One contract test is enough: UTF-8 BOM, semicolon headers, RFC-style escaping, fallback values, sanitized timestamped filename. Avoid a large E2E suite for this low-value feature.", "PR"),
            ("QA-IMP-038", "P1", "UI", "600 px panel layout", "No file, validating, error alert, long filename, history", "No horizontal clipping; actions remain reachable; error dialog fits viewport; long filenames/messages wrap; scroll is contained in intended region.", "PR"),
            ("QA-IMP-039", "P2", "UI", "Keyboard and focus", "Keyboard-only path", "Tabs, download, uploader, cancel, submit, dialog, and close controls are reachable in logical order; focus remains visible and returns from dialog.", "Release"),
            ("QA-IMP-040", "P1", "UI", "Spanish encoding regression", "All visible labels/messages", "No mojibake such as ValidaciÃ³n, botÃ³n, or mÃ¡s appears in DOM, snackbar, CSV headers, or downloaded filenames. Current source snapshot shows encoding risk.", "PR"),
        ],
    ),
]

for title, rows in suites:
    add_heading(doc, title, 2)
    add_table(doc, scenario_headers, rows, scenario_widths, font_size=7.35, priority_column=1)

doc.add_page_break()
add_heading(doc, "5. High-risk API verification recipes", 1)

add_heading(doc, "5.1 Create from an available plate", 2)
for text in (
    "Reserve a disposable plate with GET /electrical-assets/article-resume/available-plates?search=0&page=0&size=N and filter out baja/devolución states.",
    "Build a minimal workbook using the exact downloaded header row. Populate all 7 required columns and valid relationships.",
    "Drive the UI upload and capture POST /electrical-assets/import/validate?loadTypeId=1. Assert multipart file name and successful totalRows=1.",
    "Click Importar activo and capture POST /electrical-assets/import/load. Assert query params and the same file payload.",
    "Poll GET /electrical-assets/article-resume/{plate} only for a short, bounded consistency window, then compare normalized persisted values.",
):
    add_bullet(doc, text)

add_heading(doc, "5.2 Existing plate replacement and blank clearing", 2)
add_callout(doc, "Data safety", "This test is intentionally destructive. Use a seeded plate dedicated to this suite, run serially, and restore the baseline in finally only if PATCH restoration can reproduce null/empty values exactly.", fill=AMBER, color=AMBER_TEXT)
for text in (
    "GET the complete baseline asset and select non-empty optional fields from different nested sections.",
    "Import a row that changes required and optional fields while leaving chosen optional cells blank.",
    "GET the asset again and compare a complete expected snapshot, not only the changed fields.",
    "Assert blank values against a canonical field-level expectation: null or empty string. Do not accept either interchangeably unless the API contract explicitly does.",
    "Verify fields omitted by the row are cleared/preserved according to full-replacement semantics. This is the main product decision that the frontend cannot reveal.",
):
    add_bullet(doc, text)

add_heading(doc, "5.3 Massive import", 2)
for text in (
    "Use 25-50 plates for nightly QA. A 500/4999-row test is a release/performance test and needs a resettable seed pool plus a larger time budget.",
    "Generate plates dynamically and stop early with a clear skip if the reserved pool is insufficient. Do not silently reuse consumed plates.",
    "Capture validate duration, load duration, returned counts, history records, and per-plate GET results. Avoid a brittle client-side SLA until the backend team defines one.",
    "Run with workers=1 and tag the suite as stateful. Parallel Playwright workers can consume each other's available plates.",
):
    add_bullet(doc, text)

add_heading(doc, "6. Excel fixture catalog", 1)
fixture_rows = [
    ("valid_minimal_1.xlsx", "One available plate; exactly 7 required values", "P0 create smoke"),
    ("valid_batch_25.xlsx", "25 dynamically reserved plates", "Nightly massive import"),
    ("header_only.xlsx", "Approved headers, no data", "Empty spreadsheet"),
    ("missing_required_<column>.xlsx", "Seven generated variants", "Required-field matrix"),
    ("headers_<mutation>.xlsx", "Missing/renamed/duplicate/reordered/extra/trimmed", "Structure contract"),
    ("sparse_rows.xlsx", "Valid rows separated by blank rows", "Physical row numbering"),
    ("invalid_dane_relation.xlsx", "Department/municipality/DANE mismatch", "Relationship rule"),
    ("duplicate_plate.xlsx", "Same plate twice with conflicting values", "Duplicate semantics"),
    ("format_boundaries.xlsx", "Max/max+1, dates, decimals, URLs, Unicode", "Representative field types"),
    ("update_full_existing_plate.xlsx", "Existing seeded plate, many changed values", "Full replacement"),
    ("update_clear_optional.xlsx", "Existing seeded plate, selected optional cells blank", "Null/empty clearing"),
    ("mixed_atomicity.xlsx", "One valid and one invalid row", "No partial commit"),
    ("size_10mb.xlsx / size_10mb_plus_1.xlsx", "Deterministic padded workbooks", "Local size boundary"),
]
add_table(doc, ["Fixture", "Construction", "Primary use"], fixture_rows, [3.0, 4.7, 2.3], font_size=8.2)

add_heading(doc, "Fixture-generation rule", 2)
add_body(doc, "Use SheetJS to read and assert the workbook contract. For generated negative fixtures, write a minimal workbook with the exact approved header row and only the cells needed by the scenario. If preserving dropdowns, comments, named ranges, and styling matters, copy the downloaded template and use a library/workflow that preserves those Excel features; do not assume a read/write round trip preserves all workbook metadata.")

add_heading(doc, "7. Automation skeletons", 1)
add_heading(doc, "7.1 SheetJS header check", 2)
add_code_block(
    doc,
    """import fs from 'node:fs/promises';
import * as XLSX from 'xlsx';

const downloadedPath = await download.path();
expect(downloadedPath).not.toBeNull();
const buffer = await fs.readFile(downloadedPath!);
const workbook = XLSX.read(buffer, { type: 'buffer' });
expect(workbook.SheetNames).toEqual(['Plantilla', 'Instrucciones', 'Codigos Dane']);

const rows = XLSX.utils.sheet_to_json<unknown[]>(workbook.Sheets.Plantilla, {
  header: 1, raw: true, blankrows: false,
});
const headers = rows[0].map(String);
expect(headers).toEqual(EXPECTED_IMPORT_HEADERS);
expect(headers).toHaveLength(102);
expect(new Set(headers).size).toBe(102);""",
)

add_heading(doc, "7.2 UI request capture plus persisted GET", 2)
add_code_block(
    doc,
    """const validate = page.waitForResponse(r =>
  r.url().includes('/electrical-assets/import/validate?loadTypeId=1') &&
  r.request().method() === 'POST');
await importPage.upload(filePath);
expect((await validate).ok()).toBeTruthy();

const load = page.waitForResponse(r =>
  r.url().includes('/electrical-assets/import/load') &&
  r.request().method() === 'POST');
await importPage.submit.click();
const loadResponse = await load;
expect(loadResponse.ok()).toBeTruthy();
const request = loadResponse.request();
expect(new URL(request.url()).searchParams.get('loadTypeId')).toBe('1');

await expect.poll(async () => {
  const persisted = await apiContext.get(`/electrical-assets/article-resume/${plate}`, {
    headers: { Authorization: `Bearer ${authToken}` },
  });
  return persisted.status();
}).toBe(200);""",
)

add_heading(doc, "8. Locator contract and missing test IDs", 1)
add_body(doc, "The supplied test-id reference exposes the following stable anchors. Use them when they exist in the target build, but first resolve the mismatch with the inspected HTML snapshot.")
available_ids = [
    ("importActiveTabsContainer", "Import/History tabs container"),
    ("importActiveImportTab", "Import tab"),
    ("importActiveHistoryTab", "History tab"),
    ("importActiveInstructionsContainer", "Steps 01-03"),
    ("importActiveDownloadTemplateButton", "Download template"),
    ("importActiveUploaderContainer", "Uploader"),
    ("importActiveActionsContainer", "Action buttons"),
    ("importActiveCancelButton", "Cancel"),
    ("importActiveSubmitButton", "Final import"),
]
add_table(doc, ["Available data-testid", "Purpose"], available_ids, [4.3, 5.7], font_size=8.5)

missing_rows = [
    ("importActiveValidatingAlert", "Assert validation loading without text-only selectors"),
    ("importActiveValidationErrorAlert", "Differentiate transient validation failure from file-list errors"),
    ("importActiveViewErrorsButton / ExportErrorsButton", "Open/export actions"),
    ("importActiveErrorDialog / ErrorRow", "Row, column, description assertions"),
    ("importActiveFileRow / RemoveButton / RetryButton", "Uploader state transitions"),
    ("importActiveHistoryList / HistoryItem / HistoryLoader", "Pagination and record assertions"),
    ("importActiveHistoryEmptyState", "Empty history"),
]
add_heading(doc, "Recommended additions", 2)
add_table(doc, ["Proposed data-testid", "Reason"], missing_rows, [4.3, 5.7], font_size=8.5)

add_heading(doc, "9. Product decisions needed before strict assertions", 1)
decision_rows = [
    ("Existing plate", "Does final load always perform a full replacement, or only for specific asset states?"),
    ("Blank optional cell", "Canonical persisted value per field: null, empty string, zero, or preserve existing?"),
    ("Mixed batch", "Validation/load atomicity: reject all, partial success, or per-row result?"),
    ("Headers", "Must order be exact? Are additional columns rejected or ignored?"),
    (".xls", "Is legacy BIFF actually supported by backend, or only advertised by frontend accept rules?"),
    ("Mass limit", "Maximum data rows and acceptable validation/load timing within the 10 MB cap."),
    ("History", "Should successful import refresh history immediately and should failures show feedback?"),
]
add_table(doc, ["Decision", "Question to lock down"], decision_rows, [2.0, 8.0], font_size=8.5)

add_callout(doc, "Realistic stop condition", "Do not automate guessed backend rules as flexible assertions. Mark those cases as contract-discovery tests, capture the actual response, agree the desired rule with product/backend, then make the assertion strict.", fill=AMBER, color=AMBER_TEXT)

doc.add_page_break()
add_heading(doc, "Appendix A. Current 102-column header contract", 1)
add_body(doc, "Exact names below come from the supplied workbook. Required/optional classification follows its header fill and comments. For maintainability, keep the canonical list in a version-controlled JSON/TypeScript fixture and review changes intentionally.")

header_rows = []
for index, header in enumerate(plantilla["headers"], start=1):
    requirement = "Required" if header.get("fill_rgb") == "FF324B48" else "Optional"
    constraint = clean_comment(header.get("comment"))
    header_rows.append((str(index), header["column"], header.get("value") or "", requirement, constraint))

add_table(
    doc,
    ["#", "Col", "Exact header", "Req.", "Workbook comment / constraint"],
    header_rows,
    [0.42, 0.48, 3.45, 0.7, 4.95],
    font_size=7.15,
)

add_heading(doc, "Appendix B. Source and implementation observations", 1)
source_rows = [
    ("Import component", r"C:\Users\marango\Desktop\MfBienesElectricos\src\app\pages\actives\pages\import-active\import-active.component.ts/.html"),
    ("Template", r"C:\Users\marango\Desktop\plantilla_importacion_2026-07-10_13.43.xlsx"),
    ("Reference maps", "Asset editing map, Create Asset robustness PDF, Import Active test-id PDF"),
    ("Existing API pattern", r"tests/editAsset/apiPayload.spec.ts and tests/editAsset/apiResponse.spec.ts"),
    ("Known frontend gaps", "Placeholder alert title 'Prueba'; no visible template/history failure; history refresh only on tab selection; possible mojibake."),
]
add_table(doc, ["Evidence", "Location / observation"], source_rows, [2.0, 8.0], font_size=8.2)

# Document metadata and update-fields hint.
doc.core_properties.title = "ImportActiveComponent - Comprehensive Test Scenario Map"
doc.core_properties.subject = "QA automation strategy for electrical asset Excel import"
doc.core_properties.author = "QA Automation"
doc.core_properties.keywords = "Playwright, Excel, SheetJS, electrical assets, import, API testing"

settings = doc.settings.element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

doc.save(OUTPUT)
print(OUTPUT)
