from __future__ import annotations

import json
import sys
from pathlib import Path

TOOLS = Path(r"C:\Users\marango\Desktop\Suite Bienes Eléctricos\.tools")
sys.path.insert(0, str(TOOLS))

import fitz  # type: ignore
from openpyxl import load_workbook  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore


ROOT = Path(r"C:\Users\marango\Desktop\Suite Bienes Eléctricos")
OUT = ROOT / "tmp" / "reference_extracts"
OUT.mkdir(parents=True, exist_ok=True)

PDFS = {
    "asset_editing_map": Path(r"C:\Users\marango\Desktop\Asset_Editing_Component_Comprehensive_Test_Scenario_Map.pdf"),
    "import_active_testids": Path(r"C:\Users\marango\Downloads\import_active_testids.pdf"),
    "active_create_robustness": Path(r"C:\Users\marango\Desktop\activeCreaterobustesness.pdf"),
}

summary: dict[str, object] = {"pdfs": {}, "xlsx": {}}

for name, path in PDFS.items():
    doc = fitz.open(path)
    text_chunks: list[str] = []
    render_dir = OUT / name
    render_dir.mkdir(parents=True, exist_ok=True)
    for index, page in enumerate(doc):
        text_chunks.append(f"\n===== PAGE {index + 1} =====\n{page.get_text('text')}")
        pix = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
        pix.save(render_dir / f"page-{index + 1:02d}.png")
    text = "".join(text_chunks)
    (OUT / f"{name}.txt").write_text(text, encoding="utf-8")
    summary["pdfs"][name] = {
        "path": str(path),
        "pages": doc.page_count,
        "characters": len(text),
        "render_dir": str(render_dir),
    }
    doc.close()

xlsx_path = Path(r"C:\Users\marango\Desktop\plantilla_importacion_2026-07-10_13.43.xlsx")
workbook = load_workbook(xlsx_path, data_only=False, read_only=False)
xlsx_summary: dict[str, object] = {
    "path": str(xlsx_path),
    "sheets": [],
    "defined_names": [name for name in workbook.defined_names],
}

for sheet in workbook.worksheets:
    rows = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True):
        rows.append([value for value in row])
    merged = [str(cell_range) for cell_range in sheet.merged_cells.ranges]
    validations = []
    if sheet.data_validations is not None:
        for validation in sheet.data_validations.dataValidation:
            validations.append(
                {
                    "type": validation.type,
                    "formula1": validation.formula1,
                    "formula2": validation.formula2,
                    "sqref": str(validation.sqref),
                    "allow_blank": validation.allow_blank,
                }
            )
    xlsx_summary["sheets"].append(
        {
            "title": sheet.title,
            "dimensions": sheet.calculate_dimension(),
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_cells": merged,
            "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "auto_filter": str(sheet.auto_filter.ref) if sheet.auto_filter.ref else None,
            "sheet_state": sheet.sheet_state,
            "sample_rows": rows,
            "data_validations": validations,
            "headers": [
                {
                    "column": get_column_letter(index),
                    "value": cell.value,
                    "fill_type": cell.fill.fill_type,
                    "fill_rgb": cell.fill.fgColor.rgb,
                    "fill_indexed": cell.fill.fgColor.indexed,
                    "font_color_type": cell.font.color.type if cell.font.color else None,
                    "font_color_rgb": cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None,
                    "bold": cell.font.bold,
                    "comment": cell.comment.text if cell.comment else None,
                    "hidden": sheet.column_dimensions[get_column_letter(index)].hidden,
                    "width": sheet.column_dimensions[get_column_letter(index)].width,
                }
                for index, cell in enumerate(sheet[1], start=1)
                if cell.__class__.__name__ != "MergedCell"
            ],
        }
    )

summary["xlsx"] = xlsx_summary
(OUT / "reference_summary.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2, default=str),
    encoding="utf-8",
)
print(
    json.dumps(
        {
            "pdf_pages": {name: data["pages"] for name, data in summary["pdfs"].items()},
            "xlsx_sheets": [
                {
                    "title": sheet["title"],
                    "dimensions": sheet["dimensions"],
                    "headers": len(sheet["headers"]),
                }
                for sheet in xlsx_summary["sheets"]
            ],
            "summary_file": str(OUT / "reference_summary.json"),
        },
        ensure_ascii=False,
        indent=2,
    )
)
