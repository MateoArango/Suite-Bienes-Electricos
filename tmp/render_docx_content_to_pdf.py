from __future__ import annotations

import html
import sys
from pathlib import Path

TOOLS = Path(r"C:\Users\marango\Desktop\Suite Bienes Eléctricos\.tools")
sys.path.insert(0, str(TOOLS))

from docx import Document  # type: ignore
from docx.oxml.ns import qn  # type: ignore
from docx.table import Table as DocxTable  # type: ignore
from docx.text.paragraph import Paragraph as DocxParagraph  # type: ignore
from reportlab.lib import colors  # type: ignore
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT  # type: ignore
from reportlab.lib.pagesizes import letter, landscape  # type: ignore
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore
from reportlab.lib.units import inch  # type: ignore
from reportlab.platypus import (  # type: ignore
    BaseDocTemplate,
    Frame,
    LongTable,
    PageBreak,
    PageTemplate,
    Paragraph,
    Preformatted,
    Spacer,
    TableStyle,
)


ROOT = Path(r"C:\Users\marango\Desktop\Suite Bienes Eléctricos")
DOCX = ROOT / "output" / "docx" / "ImportActiveComponent_Automation_Test_Scenario_Map.docx"
PDF = ROOT / "output" / "pdf" / "ImportActiveComponent_Automation_Test_Scenario_Map.pdf"
PDF.parent.mkdir(parents=True, exist_ok=True)

NAVY = colors.HexColor("#17324D")
BLUE = colors.HexColor("#2E74B5")
DARK_BLUE = colors.HexColor("#1F4D78")
MUTED = colors.HexColor("#5B6573")
BLACK = colors.HexColor("#111827")
BORDER = colors.HexColor("#AEB7C2")


def block_items(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


def cell_fill(cell):
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    value = shd.get(qn("w:fill")) if shd is not None else None
    if value and value not in ("auto", "none"):
        try:
            return colors.HexColor(f"#{value}")
        except Exception:
            return colors.white
    return colors.white


def grid_widths(table, available_width):
    cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    raw = [int(col.get(qn("w:w"), "0")) for col in cols]
    if len(raw) != len(table.columns) or sum(raw) <= 0:
        return [available_width / len(table.columns)] * len(table.columns)
    scale = available_width / sum(raw)
    return [value * scale for value in raw]


def esc(text):
    return html.escape(text or "").replace("\n", "<br/>")


styles = getSampleStyleSheet()
body = ParagraphStyle(
    "Body",
    parent=styles["BodyText"],
    fontName="Helvetica",
    fontSize=8.4,
    leading=10.2,
    textColor=BLACK,
    spaceAfter=4,
)
h1 = ParagraphStyle("H1", parent=body, fontName="Helvetica-Bold", fontSize=14.5, leading=17, textColor=BLUE, spaceBefore=9, spaceAfter=6, keepWithNext=True)
h2 = ParagraphStyle("H2", parent=body, fontName="Helvetica-Bold", fontSize=11.7, leading=14, textColor=BLUE, spaceBefore=7, spaceAfter=5, keepWithNext=True)
h3 = ParagraphStyle("H3", parent=body, fontName="Helvetica-Bold", fontSize=10.2, leading=12, textColor=DARK_BLUE, spaceBefore=6, spaceAfter=4, keepWithNext=True)
title = ParagraphStyle("Title", parent=body, fontName="Helvetica-Bold", fontSize=22, leading=25, textColor=NAVY, spaceAfter=3)
subtitle = ParagraphStyle("Subtitle", parent=body, fontName="Helvetica", fontSize=11.5, leading=14, textColor=MUTED, spaceAfter=8)
kicker = ParagraphStyle("Kicker", parent=body, fontName="Helvetica-Bold", fontSize=9.4, leading=11, textColor=BLUE, spaceAfter=2)
bullet = ParagraphStyle("Bullet", parent=body, leftIndent=16, firstLineIndent=-8, bulletIndent=4, spaceAfter=3)
code_style = ParagraphStyle("Code", parent=body, fontName="Courier", fontSize=6.7, leading=8.1, textColor=colors.HexColor("#263445"), spaceAfter=0)

page_w, page_h = landscape(letter)
margin = 0.5 * inch
available_w = page_w - 2 * margin


def page_decor(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica-Bold", 7.4)
    canvas.setFillColor(MUTED)
    canvas.drawString(margin, page_h - 0.28 * inch, "ImportActiveComponent")
    canvas.setFont("Helvetica", 7.4)
    canvas.drawRightString(page_w - margin, page_h - 0.28 * inch, "QA automation scenario map | 2026-07-14")
    canvas.drawRightString(page_w - margin, 0.24 * inch, f"Electrical Assets Import | Page {doc.page}")
    canvas.restoreState()


pdf_doc = BaseDocTemplate(
    str(PDF),
    pagesize=landscape(letter),
    leftMargin=margin,
    rightMargin=margin,
    topMargin=0.44 * inch,
    bottomMargin=0.40 * inch,
    title="ImportActiveComponent - Comprehensive Test Scenario Map",
    author="QA Automation",
)
frame = Frame(pdf_doc.leftMargin, pdf_doc.bottomMargin, pdf_doc.width, pdf_doc.height, id="normal")
pdf_doc.addPageTemplates(PageTemplate(id="all", frames=[frame], onPage=page_decor))

source = Document(DOCX)
story = []

for block in block_items(source):
    if isinstance(block, DocxParagraph):
        xml = block._p.xml
        if 'w:type="page"' in xml:
            story.append(PageBreak())
        text = block.text.strip()
        if not text:
            continue
        style_name = block.style.name if block.style is not None else "Normal"
        if style_name == "Heading 1":
            style = h1
        elif style_name == "Heading 2":
            style = h2
        elif style_name == "Heading 3":
            style = h3
        elif style_name.startswith("List Bullet"):
            style = bullet
            text = f"- {text}"
        elif text == "ImportActiveComponent - Comprehensive Test Scenario Map":
            style = title
        elif text == "Playwright + APIRequestContext + Excel contract checks with SheetJS":
            style = subtitle
        elif text == "QA AUTOMATION TEST STRATEGY":
            style = kicker
        else:
            style = body
        story.append(Paragraph(esc(text), style))
    else:
        row_count = len(block.rows)
        col_count = len(block.columns)
        widths = grid_widths(block, available_w)
        data = []
        fills = []
        single_cell_text = block.cell(0, 0).text.lstrip() if row_count == 1 and col_count == 1 else ""
        is_code = row_count == 1 and col_count == 1 and (
            single_cell_text.startswith("import ") or single_cell_text.startswith("const validate")
        )
        for row_index, row in enumerate(block.rows):
            values = []
            fill_row = []
            for cell in row.cells:
                text = "\n".join(p.text for p in cell.paragraphs).strip()
                if is_code:
                    values.append(Preformatted(text, code_style))
                else:
                    font_size = 6.25 if col_count >= 7 else 6.75 if col_count >= 5 else 7.4 if col_count >= 3 else 8.0
                    cell_style = ParagraphStyle(
                        f"cell_{col_count}_{row_index}",
                        parent=body,
                        fontName="Helvetica-Bold" if row_index == 0 and row_count > 1 else "Helvetica",
                        fontSize=font_size,
                        leading=font_size + 1.4,
                        alignment=TA_CENTER if col_count >= 5 and len(text) < 18 else TA_LEFT,
                        textColor=NAVY if row_index == 0 and row_count > 1 else BLACK,
                    )
                    values.append(Paragraph(esc(text), cell_style))
                fill_row.append(cell_fill(cell))
            data.append(values)
            fills.append(fill_row)

        table = LongTable(
            data,
            colWidths=widths,
            repeatRows=1 if row_count > 1 else 0,
            splitByRow=1,
            hAlign="LEFT",
        )
        commands = [
            ("GRID", (0, 0), (-1, -1), 0.35, BORDER),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for row_index, row in enumerate(fills):
            for col_index, fill in enumerate(row):
                if fill != colors.white:
                    commands.append(("BACKGROUND", (col_index, row_index), (col_index, row_index), fill))
        table.setStyle(TableStyle(commands))
        story.append(table)
        story.append(Spacer(1, 4))

pdf_doc.build(story)
print(PDF)
